#!/usr/bin/env python3
"""Independent check of a workbook: LibreOffice recalculation, error cells, banned functions, macros. Short output.
usage: verify_workbook.py path/to/workbook.xlsx     (needs LibreOffice `soffice` on the PATH and the openpyxl package)
Exit code 1 if there are error cells or banned functions. A macro project is reported (macros=True) but does not change the exit code."""
import sys, os, re, subprocess, shutil, tempfile, hashlib, zipfile
if len(sys.argv) != 2 or not os.path.exists(sys.argv[1]):
    sys.exit("usage: verify_workbook.py path/to/workbook.xlsx")
if not shutil.which("soffice"):
    sys.exit("soffice (LibreOffice) was not found on the PATH")
try:
    import openpyxl
except ImportError:
    sys.exit("the openpyxl package is needed: pip install openpyxl")
src = sys.argv[1]
ERR = ('#REF!', '#NAME?', '#VALUE!', '#N/A', '#DIV/0!', '#NUM!', '#NULL!', 'Err:')
ban = re.compile(r'\b(XLOOKUP|FILTER|SORT|UNIQUE|LET|LAMBDA|MAXIFS|MINIFS|IFS|SWITCH|TEXTJOIN|CONCAT|SEQUENCE|XMATCH)\s*\(', re.I)


def check(td):
    shutil.copy2(src, f"{td}/in.xlsx")
    try:
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{td}/lo", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", f"{td}/out", f"{td}/in.xlsx"], capture_output=True, timeout=240)
    except subprocess.TimeoutExpired:
        sys.exit("LibreOffice did not finish within 240 seconds")
    if not os.path.exists(f"{td}/out/in.xlsx"):
        sys.exit("LibreOffice did not produce a recalculated copy of the workbook")
    wf = openpyxl.load_workbook(f"{td}/in.xlsx"); wv = openpyxl.load_workbook(f"{td}/out/in.xlsx", data_only=True)
    errs = bans = forms = 0
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    forms += 1; bans += bool(ban.search(c.value))
        for row in wv[ws.title].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERR): errs += 1
    sha = hashlib.sha256(open(src, 'rb').read()).hexdigest()[:16]
    macros = any(n.lower().endswith("vbaproject.bin") for n in zipfile.ZipFile(src).namelist())
    print(f"file={os.path.basename(src)} bytes={os.path.getsize(src)} sha256={sha} sheets={len(wf.sheetnames)} formulas={forms} error_cells={errs} banned_functions={bans} macros={macros}")
    print("sheets:", ", ".join(wf.sheetnames))
    return 1 if (errs or bans) else 0


td = tempfile.mkdtemp(prefix="wbv_")
try:
    code = check(td)
finally:
    shutil.rmtree(td, ignore_errors=True)
sys.exit(code)
